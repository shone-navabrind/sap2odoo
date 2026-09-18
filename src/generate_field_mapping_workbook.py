"""
Generates SAP_Field_Mapping.xlsx - one worksheet per sheet-object in src/registry.py, listing:
  - every real SAP field that actually came back for that object (from output_raw/*.json's
    fields_present, not a guessed schema), grouped by which raw entity it came from
  - a sample real value pulled from the raw data (evidence, not a description guess)
  - the Odoo field it's mapped to (transcribed from the actual transform_odoo.py logic for that
    object), or "(not mapped)" if the field exists in SAP but isn't currently used
  - the target Odoo model and a note

For objects still pending_mapping or not_in_bydesign (no raw data exists), the sheet has a
single info row instead - there's nothing to list fields for yet.

Run: python -m src.generate_field_mapping_workbook
"""

import json
import logging
import os
import re
import sys

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.registry import REGISTRY

logger = logging.getLogger("sap2odoo.generate_field_mapping_workbook")

RAW_DIR = "output_raw"
OUTPUT_PATH = "SAP_Field_Mapping.xlsx"

# Transcribed directly from src/transform_odoo.py's actual field usage per object (not
# guessed) - {sheet_no: {sap_field_name: (odoo_field, odoo_model, note)}}. Fields present in
# the raw data but NOT in this dict are real SAP fields that simply aren't mapped yet - the
# sheet shows those as "(not mapped)" rather than omitting them.
FIELD_MAP = {
    5: {  # Banks
        "CBANK_NAME": ("name", "res.bank", "From RPBUPCSD/RPBUPSPP"),
        "CBANK_NAT_COUNTRY": ("country_id/id", "res.bank", ""),
        "CBANK_ACCOUNT_ID": ("acc_number", "res.partner.bank", "Used if CBANK_IBAN is blank"),
        "CBANK_IBAN": ("acc_number", "res.partner.bank", "Preferred over CBANK_ACCOUNT_ID"),
        "CBP_UUID": ("partner_id/id", "res.partner.bank", ""),
    },
    6: {  # Cost Centers
        "UUID": ("id (external ID key)", "account.analytic.account", ""),
        "ID": ("code", "account.analytic.account", ""),
        "MostRecentDefaultName": ("name", "account.analytic.account", ""),
    },
    4: {  # Payment Terms
        "PaymentTermsCode": ("id (external ID key)", "account.payment.term", "khcustomerinvoice side"),
        "PaymentTermsCodeText": ("name", "account.payment.term", "khcustomerinvoice side"),
        "Code": ("id (external ID key)", "account.payment.term", "khsupplierinvoice side"),
        "CodeText": ("name", "account.payment.term", "khsupplierinvoice side"),
    },
    9: {},  # Open Vendor Bills - same mapping as #51 (Vendor Bills), filtered by LifeCycleStatusCode
    10: {},  # Customer Payments - same mapping as #11 (shared PaymentCollection, split by partner_type)
    11: {  # Vendor Payments (Customer Payments #10 shares this exact map)
        "ObjectID": ("id (external ID key)", "account.payment", ""),
        "BusinessPartnerID": ("partner_id/id", "account.payment", "Also used to look up customer_rank/supplier_rank in res_partner.csv to derive payment_type/partner_type"),
        "TransactionCurrencyAmount": ("amount", "account.payment", ""),
        "AccountingTransactionDate": ("date", "account.payment", ""),
        "TransactionCurrencyCode": ("currency_id/id", "account.payment", ""),
    },
    12: {  # Bank Statements
        "ObjectID": ("id (external ID key)", "account.bank.statement", ""),
        "ID": ("name", "account.bank.statement", ""),
        "Date": ("date", "account.bank.statement", ""),
        "OpeningBalanceAmount": ("balance_start", "account.bank.statement", ""),
        "ClosingBalanceAmount": ("balance_end_real", "account.bank.statement", ""),
    },
    17: {  # Customers (Vendors #46 shares this exact map)
        "CBP_UUID": ("id (external ID key)", "res.partner", ""),
        "TBP_UUID": ("name", "res.partner", "This field is the partner's TEXT/name despite the field name"),
        "CSTREET_NAME": ("street", "res.partner", ""),
        "CCITY_NAME": ("city", "res.partner", ""),
        "CSTREET_POSTAL": ("zip", "res.partner", ""),
        "CCOUNTRY_CODE": ("country_id/id", "res.partner", ""),
        "CPHONE_NR": ("phone", "res.partner", ""),
        "CEMAIL_URI": ("email", "res.partner", ""),
        "CWEB_URI": ("website", "res.partner", ""),
        "CDELIVERY_BLOCK": ("active (inverted)", "res.partner", ""),
        "CTAX_NUMBER": ("vat", "res.partner", "From RPBUPATAXNUMBERS, joined by CBUPA_UUID/CBP_UUID"),
        "CBUPA_UUID": ("(join key to CBP_UUID)", "res.partner", "RPBUPATAXNUMBERS side"),
    },
    46: {},  # Vendors - identical mapping to #17 Customers
    18: {  # Salespersons
        "ObjectID": ("id (external ID key)", "res.users", ""),
        "FormattedName": ("name", "res.users", ""),
        "Email": ("email / login", "res.users", "From WorkplaceAddressCollection, joined by ParentObjectID"),
        "Phone": ("phone", "res.users", "From WorkplaceAddressCollection"),
    },
    20: {  # Opportunities (Open #21 / Closed #22 share this map, split by LifeCycleStatusCode)
        "ObjectID": ("id (external ID key)", "crm.lead", ""),
        "Description": ("name", "crm.lead", ""),
        "ExpectedRevenueAmount": ("expected_revenue", "crm.lead", ""),
        "ChanceOfSuccessPercent": ("probability", "crm.lead", ""),
        "LifeCycleStatusCode": ("(filter only, not mapped)", "crm.lead", "1/2=Open, 4/5=Won/Lost - used to split crm_lead_open.csv / crm_lead_closed.csv"),
    },
    21: {},  # Open Opportunities - same as #20
    22: {},  # Closed Opportunities - same as #20
    27: {  # Warehouses (subset of #28 Locations, filtered)
        "ObjectID": ("id (external ID key)", "stock.warehouse", ""),
        "Name": ("name", "stock.warehouse", ""),
        "ID": ("code", "stock.warehouse", ""),
        "InventoryManagedLocationIndicator": ("(filter only, not mapped)", "stock.warehouse", "True = this location becomes a warehouse row"),
    },
    28: {  # Locations
        "ObjectID": ("id (external ID key)", "stock.location", ""),
        "Name": ("name", "stock.location", ""),
    },
    29: {  # UOM
        "Code": ("id (external ID key)", "uom.uom", ""),
        "Description": ("name", "uom.uom", ""),
    },
    42: {  # Work Centers
        "ResourceID": ("id (external ID key) / code", "mrp.workcenter", "From khproductionorder's OperationCollection"),
        "ResourceDescription": ("name", "mrp.workcenter", ""),
    },
    44: {  # Manufacturing Orders
        "ObjectID": ("id (external ID key)", "mrp.production", ""),
        "ID": ("name", "mrp.production", ""),
        "RequestedStartDateTime": ("date_planned_start", "mrp.production", ""),
        "RequestedEndDateTime": ("date_planned_finished", "mrp.production", ""),
        "BillOfMaterialID": ("(not mapped)", "mrp.production", "Referenced but no component/qty data attached on this tenant - see CLAUDE.md"),
    },
    50: {  # Purchase Orders
        "ObjectID": ("id (external ID key)", "purchase.order / purchase.order.line", ""),
        "ID": ("name", "purchase.order", ""),
        "CreationDateTime": ("date_order", "purchase.order", ""),
        "CurrencyCode": ("currency_id/id", "purchase.order", ""),
        "TotalNetAmount": ("amount_total", "purchase.order", ""),
        "PartyID": ("partner_id/id", "purchase.order", "From SupplierCollection, joined by ParentObjectID"),
        "ParentObjectID": ("order_id/id", "purchase.order.line", "ItemCollection side"),
        "ProductID": ("product_id/id", "purchase.order.line", ""),
        "Description": ("name", "purchase.order.line", ""),
        "Quantity": ("product_qty", "purchase.order.line", ""),
        "NetUnitPriceAmount": ("price_unit", "purchase.order.line", ""),
    },
    51: {  # Vendor Bills (Open Vendor Bills #9 shares this map, filtered by LifeCycleStatusCode)
        "ObjectID": ("id (external ID key)", "account.move / account.move.line", ""),
        "InvoiceDate": ("invoice_date", "account.move", ""),
        "TotalNetAmountCurrencyCode": ("currency_id/id", "account.move", ""),
        "TotalNetAmount": ("amount_total", "account.move", ""),
        "LifeCycleStatusCode": ("(filter only, not mapped)", "account.move", "Excludes Paid=12/Canceled=9/Voided=7 for account_move_open_vendor.csv"),
        "PartyID": ("partner_id/id", "account.move", "From SellerPartyCollection, joined by ParentObjectID"),
        "ParentObjectID": ("move_id/id", "account.move.line", "ItemCollection side"),
        "ProductID": ("product_id/id", "account.move.line", ""),
        "Description": ("name", "account.move.line", ""),
        "Quantity": ("quantity", "account.move.line", ""),
        "NetUnitPriceAmount": ("price_unit", "account.move.line", ""),
    },
    57: {  # Products
        "ObjectID": ("id (external ID key)", "product.template", ""),
        "InternalID": ("default_code", "product.template", "Also used as the id if present"),
        "Description": ("name", "product.template", ""),
        "UUID": ("(join key)", "product.template", "Links to vmumaterialvaluationdata's MaterialUUID"),
        "BaseMeasureUnitCode": ("uom_id/id", "product.template", ""),
        "Text": ("description", "product.template", "TextCollection where TypeCode=10006 (Detailed Description)"),
        "TypeCode": ("(filter only, not mapped)", "product.template", "TextCollection - only TypeCode=10006 rows are used"),
        "PurchasingMeasureUnitCode": ("uom_po_id/id", "product.template", "PurchasingCollection; presence also sets purchase_ok=True"),
        "SalesMeasureUnitCode": ("(presence signal, not mapped)", "product.template", "SalesCollection; presence sets sale_ok=True"),
        "ProductCategoryInternalID": ("categ_id/id", "product.template", "ProductCategoryCollection"),
        "MaterialUUID": ("(join key)", "product.template", "vmumaterialvaluationdata's MaterialValuationDataCollection"),
        "Amount": ("standard_price", "product.template", "ValuationPriceCollection, latest by StartDate"),
        "StartDate": ("(used to pick latest price, not mapped)", "product.template", "ValuationPriceCollection"),
    },
    58: {  # Product Categories
        "ProductCategoryInternalID": ("id (external ID key)", "product.category", ""),
        "Description": ("name", "product.category", ""),
    },
    59: {  # Pricelists (Discount Rules #60 / Customer Price Lists #61 share this source, not separately mapped)
        "ObjectID": ("id (external ID key)", "product.pricelist", ""),
        "CurrencyCode": ("currency_id/id", "product.pricelist", ""),
        "CustomerUUID": ("(not mapped)", "product.pricelist", "Hyphenated GUID doesn't match res_partner's numeric-ID external IDs - known limitation, see CLAUDE.md"),
    },
    63: {  # Sales Orders
        "ObjectID": ("id (external ID key)", "sale.order / sale.order.line", ""),
        "ID": ("name", "sale.order", ""),
        "PostingDateTime": ("date_order", "sale.order", ""),
        "NetAmountCurrencyCode": ("currency_id/id", "sale.order", ""),
        "NetAmount": ("amount_total / price_subtotal", "sale.order / sale.order.line", "Header uses SalesOrderCollection's NetAmount; line uses ItemCollection's NetAmount"),
        "PartyID": ("partner_id/id", "sale.order", "From BuyerPartyCollection - see resolve_party() heuristic in CLAUDE.md"),
        "ParentObjectID": ("order_id/id", "sale.order.line", "ItemCollection side"),
        "ProductID": ("product_id/id", "sale.order.line", "From ItemProductCollection, joined by ParentObjectID=Item.ObjectID"),
        "Description": ("name", "sale.order.line", ""),
    },
    64: {  # Deliveries
        "ObjectID": ("id (external ID key)", "stock.picking / stock.picking.line", ""),
        "ID": ("name", "stock.picking", ""),
        "CreationDateTime": ("scheduled_date", "stock.picking", ""),
        "PartyID": ("partner_id/id", "stock.picking", "From BuyerPartyCollection"),
        "ParentObjectID": ("picking_id/id", "stock.picking.line", "ItemCollection side"),
        "ProductID": ("product_id/id / name", "stock.picking.line", ""),
    },
    65: {  # Customer Invoices
        "ObjectID": ("id (external ID key)", "account.move / account.move.line", ""),
        "Date": ("invoice_date", "account.move", ""),
        "TotalNetAmountCurrencyCode": ("currency_id/id", "account.move", ""),
        "TotalNetAmount": ("amount_total", "account.move", ""),
        "PartyID": ("partner_id/id", "account.move", "From BuyerPartyCollection"),
        "ParentObjectID": ("move_id/id", "account.move.line", "ItemCollection side"),
        "ProductID": ("product_id/id", "account.move.line", ""),
        "Description": ("name", "account.move.line", ""),
        "Quantity": ("quantity", "account.move.line", ""),
    },
}

# Objects whose FIELD_MAP is empty because they share another object's mapping 1:1 (same raw
# source, same transform, just filtered or split differently).
SHARES_MAPPING_OF = {9: 51, 10: 11, 21: 20, 22: 20, 46: 17}

HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
MAPPED_FILL = PatternFill(start_color="D9EAD3", end_color="D9EAD3", fill_type="solid")
UNMAPPED_FILL = PatternFill(start_color="FCE5CD", end_color="FCE5CD", fill_type="solid")


def _safe_sheet_name(name, sheet_no):
    cleaned = re.sub(r'[\\/*?:\[\]]', "-", name)[:25]
    return f"{sheet_no}-{cleaned}" if sheet_no else f"X-{cleaned}"


def _load_raw_entities(raw_sources):
    """Returns [(entity_set_name, fields_present, rows), ...] for an object's raw sources."""
    entities = []
    for filename in raw_sources:
        path = os.path.join(RAW_DIR, filename)
        if not os.path.isfile(path):
            continue
        with open(path, encoding="utf-8") as f:
            payload = json.load(f)
        entities.append((payload.get("entity_set", filename), payload.get("fields_present", []), payload.get("rows", [])))
    return entities


def _sample_value(rows, field):
    for row in rows:
        value = row.get(field)
        if value not in (None, "", "false", False):
            text = str(value)
            return text[:60] + ("..." if len(text) > 60 else "")
    return ""


def _autofit(ws):
    for col_cells in ws.columns:
        length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=10)
        ws.column_dimensions[get_column_letter(col_cells[0].column)].width = min(max(length + 2, 12), 60)


def _build_index_sheet(wb, sheet_name_by_sheet_no):
    ws = wb.create_sheet("Index", 0)
    header = ["Sheet #", "Object Name", "Module", "Category", "Mandatory", "Status", "Odoo Model", "Go to tab"]
    ws.append(header)
    for col in range(1, len(header) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT

    for obj in REGISTRY:
        tab_name = sheet_name_by_sheet_no[obj.sheet_no]
        ws.append([
            obj.sheet_no if obj.sheet_no else "(added)",
            obj.name,
            obj.module,
            obj.category,
            "Yes" if obj.mandatory else "Optional",
            obj.status,
            obj.odoo_model,
            tab_name,
        ])
        link_cell = ws.cell(row=ws.max_row, column=8)
        link_cell.hyperlink = f"#'{tab_name}'!A1"
        link_cell.font = Font(color="0563C1", underline="single")
        status_cell = ws.cell(row=ws.max_row, column=6)
        status_cell.fill = MAPPED_FILL if obj.status == "built" else (
            UNMAPPED_FILL if obj.status == "pending_mapping" else PatternFill()
        )
    ws.freeze_panes = "A2"
    _autofit(ws)


def build_workbook():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    used_names = set()
    sheet_name_by_sheet_no = {}

    for obj in REGISTRY:
        sheet_name = _safe_sheet_name(obj.name, obj.sheet_no)
        base_name = sheet_name
        i = 2
        while sheet_name in used_names:
            sheet_name = f"{base_name[:28]}-{i}"
            i += 1
        used_names.add(sheet_name)
        sheet_name_by_sheet_no[obj.sheet_no] = sheet_name
        ws = wb.create_sheet(sheet_name)

        ws.append([f"#{obj.sheet_no} - {obj.name}"])
        ws["A1"].font = Font(bold=True, size=14)
        ws.append([
            f"Module: {obj.module}  |  Category: {obj.category}  |  "
            f"Mandatory: {'Yes' if obj.mandatory else 'Optional'}  |  Odoo model: {obj.odoo_model}  |  "
            f"Status: {obj.status.upper()}"
        ])
        ws["A2"].font = Font(italic=True)
        ws.append([obj.note])
        ws["A3"].font = Font(size=9, color="666666")
        ws.append([])

        if obj.status != "built" or not obj.raw_sources:
            ws.append(["No SAP data extracted yet for this object - see the notes above and CLAUDE.md."])
            ws["A5"].font = Font(italic=True, color="CC0000")
            _autofit(ws)
            continue

        field_map_key = SHARES_MAPPING_OF.get(obj.sheet_no, obj.sheet_no)
        field_map = FIELD_MAP.get(field_map_key, {})

        header_row = ["SAP Entity (raw source)", "SAP Field Name", "Sample Value", "Mapped Odoo Field", "Target Odoo Model", "Notes"]
        ws.append(header_row)
        header_idx = ws.max_row
        for col in range(1, len(header_row) + 1):
            cell = ws.cell(row=header_idx, column=col)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT

        entities = _load_raw_entities(obj.raw_sources)
        for entity_name, fields, rows in entities:
            for field in sorted(fields):
                mapping = field_map.get(field)
                odoo_field = mapping[0] if mapping else "(not mapped)"
                odoo_model = mapping[1] if mapping else obj.odoo_model
                note = mapping[2] if mapping else ""
                ws.append([entity_name, field, _sample_value(rows, field), odoo_field, odoo_model, note])
                fill = UNMAPPED_FILL if odoo_field == "(not mapped)" else MAPPED_FILL
                for col in range(1, len(header_row) + 1):
                    ws.cell(row=ws.max_row, column=col).fill = fill

        ws.freeze_panes = f"A{header_idx + 1}"
        _autofit(ws)

    _build_index_sheet(wb, sheet_name_by_sheet_no)
    return wb


def main():
    logging.basicConfig(level=logging.INFO, format="%(levelname)s %(message)s")
    wb = build_workbook()
    wb.save(OUTPUT_PATH)
    logger.info("Wrote %s (%d sheets)", OUTPUT_PATH, len(wb.sheetnames))


if __name__ == "__main__":
    sys.exit(main())
